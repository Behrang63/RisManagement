# =============================================================================
# 1. IMPORTS & DEPENDENCIES
# =============================================================================
import streamlit as st
import pandas as pd
import json
import os
import requests
import re
import plotly.express as px

# Optional Document Processing Modules
try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# =============================================================================
# 2. ENVIRONMENT BOOTSTRAPPING & DEPLOYMENT
# =============================================================================
def generate_launcher():
    """Generates an idempotent .bat file for autonomous deployment and venv detection."""
    script_path = os.path.abspath(__file__)
    base_name = os.path.splitext(os.path.basename(script_path))[0]
    bat_path = os.path.join(os.path.dirname(script_path), f"{base_name}_launcher.bat")
    
    if not os.path.exists(bat_path):
        bat_content = f"""@echo off
echo ===================================================
echo Bootstrapping Agentic Environment...
echo ===================================================
IF EXIST venv\\Scripts\\activate.bat (
    call venv\\Scripts\\activate.bat
    echo [INFO] Virtual Environment 'venv' activated.
) ELSE IF EXIST .venv\\Scripts\\activate.bat (
    call .venv\\Scripts\\activate.bat
    echo [INFO] Virtual Environment '.venv' activated.
) ELSE (
    echo [WARNING] No virtual environment found. Running globally.
)
streamlit run "{script_path}"
pause
"""
        try:
            with open(bat_path, "w", encoding="utf-8") as f:
                f.write(bat_content)
        except Exception:
            pass 

generate_launcher()

# =============================================================================
# 3. CoALA MEMORY LOADERS (Semantic Memory)
# =============================================================================
def load_semantic_memory(filename="system_rules.md"):
    """Loads static business logic and constraints into semantic memory."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    skill_dir = os.path.join(script_dir, "Ethic_SKILL")
    os.makedirs(skill_dir, exist_ok=True)
    
    file_path = os.path.join(skill_dir, filename)
    
    if not os.path.exists(file_path):
        default_rules = """# Agentic Medical Risk Evaluator - Core Semantic Rules
1. **Safety First**: Prioritize patient privacy and informed consent over operational speed.
2. **Bias Detection**: Explicitly evaluate clinical trials for systematic sampling and selection bias.
3. **Rigorous FMEA**: Assess severity (S), occurrence (O), and detection (D) metrics strictly on a 1-10 scale.
4. **Factual Output**: The output must be strictly analytical, objective, and devoid of speculative storytelling.
"""
        try:
            with open(file_path, "w", encoding="utf-8") as f:
                f.write(default_rules)
        except Exception:
            pass
        return default_rules
    else:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

# =============================================================================
# 4. CORE AGENTIC LOGIC (Maker-Checker Architecture)
# =============================================================================
def query_ollama(url: str, model: str, prompt: str, system_prompt: str = "") -> str:
    """Raw communication channel to the LLM. Unlimited timeout for local inference."""
    payload = {"model": model, "prompt": prompt, "stream": False}
    if system_prompt:
        payload["system"] = system_prompt
    try:
        response = requests.post(url, json=payload, timeout=None)
        response.raise_for_status()
        return response.json().get("response", "")
    except requests.exceptions.RequestException as e:
        return f"ERROR: Network or connection failure. Is Ollama running? Details: {str(e)}"
    except Exception as e:
        return f"ERROR: {str(e)}"

def maker_agent(url: str, model: str, document_text: str, semantic_rules: str, feedback: str = "", previous_output: str = "") -> str:
    """The Generation Engine: Creates the initial output or refines it based on stateful feedback."""
    
    system_instruction = (
        "You are the Maker Agent: an expert clinical risk auditor. "
        f"Strictly adhere to these Semantic Rules:\n{semantic_rules}\n\n"
        "OUTPUT FORMAT MANDATE: You must return ONLY a raw JSON array of up to 5 risk objects. "
        "No markdown code blocks, no conversational text, no introductions. "
        "LANGUAGE MANDATE: ALL text values MUST be in natural, grammatically correct Persian (Farsi) following standard academic and clinical writing principles. Avoid machine-translation tones and use precise medical/clinical terminology in Persian. "
        "Each object MUST contain EXACTLY these keys: 'risk' (a descriptive title of 2 to 10 words), 'description', 'S', 'O', 'D'."
    )
    
    prompt = f"DOCUMENT TO AUDIT:\n{document_text}\n\n"
    
    # Stateful Correction: Inject the failed output so the model knows what to fix
    if feedback and previous_output:
        prompt += (
            f"🚨 PREVIOUS FAILED OUTPUT (DO NOT REPEAT THIS MISTAKE):\n{previous_output}\n\n"
            f"🚨 CRITICAL EVIDENCE FROM CHECKER AGENT (FIX THESE SPECIFIC ISSUES):\n{feedback}\n\n"
            "REGENERATE VALID JSON ARRAY:"
        )
    else:
        prompt += "\n\nGENERATE FMEA RISK ANALYSIS AS JSON ARRAY:"

    return query_ollama(url, model, prompt, system_instruction)

def checker_agent(raw_output: str) -> tuple[bool, list, str]:
    """The Evaluator Engine: Strictly read-only logic and syntax audit. Never generates payload."""
    try:
        clean_text = raw_output.strip()
        
        # Robust Pre-Parse Sanitization
        json_match = re.search(r'\[\s*\{.*\}\s*\]', clean_text, re.DOTALL)
        if json_match:
            clean_text = json_match.group(0)
            
        # Clean trailing commas common in LLM outputs to prevent trivial decode errors
        clean_text = re.sub(r',\s*\}', '}', clean_text)
        clean_text = re.sub(r',\s*\]', ']', clean_text)
            
        parsed_data = json.loads(clean_text)
        
        if not isinstance(parsed_data, list):
            return False, [], "EVIDENCE: Root element is not a list. You must return a JSON array [...]."
        if len(parsed_data) == 0:
            return False, [], "EVIDENCE: JSON array is empty. At least one risk must be identified."
            
        validated_data = []
        for idx, item in enumerate(parsed_data[:5]):
            missing_keys = [k for k in ['risk', 'description', 'S', 'O', 'D'] if k not in item]
            if missing_keys:
                return False, [], f"EVIDENCE: Item at index {idx} is missing required keys: {missing_keys}."
            
            # --- AGENTIC SECURITY PATCH: ZWNJ, Placeholder & Length Check ---
            risk_title = str(item.get('risk', '')).replace('\u200c', ' ').strip()
            invalid_placeholders = ["بدون نام", "unknown", "نامشخص", "none", "null", "بدون عنوان", "بدون ریسک", "نقش", "ریسک", "خطا", "مشکل"]
            
            if not risk_title or risk_title.lower() in invalid_placeholders:
                return False, [], f"EVIDENCE: The risk title '{risk_title}' is a meaningless placeholder. You must extract a specific, highly relevant clinical risk title from the text."
                
            title_words = risk_title.split()
            # Relaxed constraint: Allows valid 2-word clinical risks (e.g., "خطای جراحی")
            if len(title_words) < 2:
                return False, [], f"EVIDENCE: The risk title '{risk_title}' is too short ({len(title_words)} word). A clinical risk title must be descriptive (at least 2 words) and grammatically correct in Persian."
            # --------------------------------------------------------
            
            try:
                s, o, d = int(item['S']), int(item['O']), int(item['D'])
                if not (1 <= s <= 10 and 1 <= o <= 10 and 1 <= d <= 10):
                    return False, [], "EVIDENCE: FMEA scores out of bounds. S,O,D must be between 1 and 10."
                
                # --- AGENTIC SECURITY PATCH: Lazy Scoring Detection ---
                if s == 5 and o == 5 and d == 5:
                    return False, [], f"EVIDENCE: Item '{item.get('risk')}' has lazy default scores (5, 5, 5). You must objectively evaluate the true severity, occurrence, and detection risk rather than defaulting to 5."
                # ------------------------------------------------------

                item['S'], item['O'], item['D'] = s, o, d
            except ValueError:
                return False, [], "EVIDENCE: FMEA scores must be pure integers."
                
            desc_words = str(item['description']).split()
            if len(desc_words) < 4:
                return False, [], f"EVIDENCE: Description for '{item.get('risk')}' is too superficial. Provide analytical depth."
                
            validated_data.append(item)
            
        return True, validated_data, "Validation Passed"

    except json.JSONDecodeError as e:
        return False, [], f"EVIDENCE: JSON parsing failed. Invalid syntax (e.g., missing quotes, trailing artifacts). Parser Error: {str(e)}"
    except Exception as e:
        return False, [], f"EVIDENCE: System exception during logic audit: {str(e)}"

def execute_agentic_loop(url: str, model: str, context_text: str, semantic_rules: str):
    """Orchestrates the Bounded Agentic Loop. Enforces Hard Boundaries & Stateful Correction."""
    MAX_RETRIES = 3
    attempt = 0
    feedback = ""
    last_raw_output = ""
    
    progress_ui = st.empty()
    status_ui = st.empty()
    
    while attempt < MAX_RETRIES:
        attempt += 1
        progress_ui.progress(attempt / MAX_RETRIES, text="در حال اجرای حلقه ایجنتیک...")
        status_ui.info(f"**🔄 تکرار {attempt}/{MAX_RETRIES}:** Maker Agent در حال تولید ساختار ریسک (ارتباط بدون محدودیت زمانی)...")
        
        # Stateful Injection: Passes `last_raw_output` so the Maker knows what to fix
        raw_output = maker_agent(url, model, context_text, semantic_rules, feedback, last_raw_output)
        
        if "ERROR" in raw_output:
            status_ui.error(f"خطای شبکه/مدل: {raw_output}")
            progress_ui.empty()
            return None # HARD BOUNDARY
            
        status_ui.warning(f"**🛡️ تکرار {attempt}/{MAX_RETRIES}:** Checker Agent در حال ارزیابی منطق و سینتکس...")
        is_valid, parsed_data, critique = checker_agent(raw_output)
        
        if is_valid:
            status_ui.success(f"✅ خروجی در تلاش {attempt} به صورت منطقی و ساختاری تایید شد.")
            progress_ui.empty()
            return parsed_data # SUCCESS BOUNDARY
        else:
            feedback = critique
            last_raw_output = raw_output # Persist output state for next iteration
            status_ui.error(f"⚠️ ارزیابی رد شد. شواهد به Maker ارسال گردید: {critique}")
            
    # GRACEFUL ESCALATION (HARD BOUNDARY)
    progress_ui.empty()
    st.error("🚨 حلقه ایجنتیک متوقف شد. مدل نتوانست در محدوده مجاز (MAX_RETRIES) خروجی معتبری تولید کند. لطفاً پرامپت را تغییر داده یا از مدل قوی‌تری استفاده کنید.")
    return None

# =============================================================================
# 5. HIGH-CRAFT CSS INJECTION (Emil Kowalski / Apple Design Principles)
# =============================================================================
def inject_premium_css():
    st.markdown("""
    <style>
        .reportview-container { direction: rtl; }
        
        /* Typography: Opt for Tabular Numbers in dashboards */
        * { font-variant-numeric: tabular-nums; font-family: 'Tahoma', sans-serif; }
        
        /* 1. Response: Sub-300ms interactions, strong ease-out cubic-bezier */
        .stButton button {
            transition: transform 160ms cubic-bezier(0.23, 1, 0.32, 1), background-color 200ms ease, border-color 200ms ease !important;
            border-radius: 8px !important;
            transform-origin: center !important;
            will-change: transform;
        }
        
        /* 2. Direct manipulation: scale(0.97) for instant active press feedback */
        .stButton button:active {
            transform: scale(0.97) !important;
            filter: blur(0.5px);
        }
        
        /* 3. Spatial Consistency: Entrance animations (Never scale(0), use scale(0.95)) */
        @keyframes subtleSlideUp {
            from { opacity: 0; transform: translateY(12px) scale(0.98); }
            to { opacity: 1; transform: translateY(0) scale(1); }
        }
        
        div[data-testid="stVerticalBlock"] > div {
            animation: subtleSlideUp 300ms cubic-bezier(0.23, 1, 0.32, 1) forwards;
        }
        
        /* 4. Materials & Depth: Translucency and Glassmorphism */
        .glass-panel {
            background: rgba(255, 255, 255, 0.03);
            backdrop-filter: blur(12px) saturate(180%);
            -webkit-backdrop-filter: blur(12px) saturate(180%);
            border: 1px solid rgba(255, 255, 255, 0.1);
            border-radius: 12px;
            padding: 24px;
            text-align: right; 
            direction: rtl;
        }
        
        .stTextArea textarea, .stTextInput input { direction: rtl; text-align: right; }
    </style>
    """, unsafe_allow_html=True)

# =============================================================================
# 6. STREAMLIT UI ORCHESTRATION
# =============================================================================
st.set_page_config(page_title="سامانه تحلیل ریسک ایجنتیک", page_icon="🛡️", layout="wide")
inject_premium_css()

# Session State Orchestration (Working Memory)
if 'ollama_url' not in st.session_state:
    st.session_state.ollama_url = "http://localhost:11434/api/generate"
if 'ollama_model' not in st.session_state:
    st.session_state.ollama_model = "llama3"
if 'verified_risks_buffer' not in st.session_state:
    st.session_state.verified_risks_buffer = []
if 'final_assessments' not in st.session_state:
    st.session_state.final_assessments = []

class DocumentProcessor:
    @staticmethod
    def extract_text(uploaded_file) -> str:
        if uploaded_file is None: return ""
        ext = os.path.splitext(uploaded_file.name)[1].lower()
        text = ""
        try:
            if ext == '.pdf' and PYPDF_AVAILABLE:
                reader = pypdf.PdfReader(uploaded_file)
                text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
            elif ext in ['.docx', '.doc'] and DOCX_AVAILABLE:
                doc = docx.Document(uploaded_file)
                text = "\n".join([para.text for para in doc.paragraphs])
            elif ext == '.txt':
                text = uploaded_file.read().decode('utf-8', errors='ignore')
            else:
                text = "فرمت فایل پشتیبانی نمی‌شود."
        except Exception as e:
            text = f"خطای استخراج: {str(e)}"
        return text.strip()

# Application Header
st.markdown("""
    <div class="glass-panel" style="margin-bottom: 24px; border-right: 6px solid #0ea5e9;">
        <h1 style="margin-top: 0; color: #0ea5e9;">🛡️ سامانه جامع تحلیل ریسک‌های پزشکی (معماری ایجنتیک)</h1>
        <p style="margin-bottom: 0; color: #64748b; font-size: 1.1rem;">
            توسعه یافته بر پایه معماری Maker-Checker، ارزیابی حلقه بسته و حافظه CoALA.
        </p>
    </div>
""", unsafe_allow_html=True)

tab_extract, tab_evaluate, tab_dashboard, tab_settings = st.tabs([
    "📥 ۱. استخراج ایجنتیک", "📋 ۲. ماتریس FMEA", "📊 ۳. داشبورد", "⚙️ ۴. تنظیمات هوش مصنوعی"
])

# --- TAB 1: Agentic Extraction ---
with tab_extract:
    col_input, col_action = st.columns([1, 1])
    
    with col_input:
        st.markdown("<div dir='rtl'><h3>📄 بارگذاری سند پروپوزال</h3></div>", unsafe_allow_html=True)
        uploaded_file = st.file_uploader("فایل خود را آپلود کنید (PDF, DOCX, TXT):", type=["pdf", "docx", "txt"])
        fallback_text = st.text_area("یا متن خام را اینجا پیست کنید:", height=150)
        
        final_text = DocumentProcessor.extract_text(uploaded_file) if uploaded_file else fallback_text

    with col_action:
        st.markdown("<div dir='rtl'><h3>🧠 کنترل اجرای حلقه</h3></div>", unsafe_allow_html=True)
        st.info("این سیستم از منطق Maker-Checker استفاده می‌کند. داده‌های تایید نشده و فِیک هرگز نمایش داده نخواهند شد.")
        
        if st.button("🚀 اجرای حلقه ارزیابی ایجنتیک (Maker-Checker)", use_container_width=True, type="primary"):
            if not final_text.strip():
                st.error("متن سند خالی است. لطفاً متنی وارد کنید.")
            else:
                semantic_rules = load_semantic_memory()
                # Run the orchestrator
                verified_data = execute_agentic_loop(
                    st.session_state.ollama_url, 
                    st.session_state.ollama_model, 
                    final_text[:2500], # Context cap
                    semantic_rules
                )
                
                if verified_data:
                    st.session_state.verified_risks_buffer = verified_data
                    st.toast("حلقه ایجنتیک با موفقیت تکمیل شد!", icon="✅")
                else:
                    st.toast("تایید داده‌ها در حلقه ایجنتیک با شکست مواجه شد.", icon="🚨")

    # Display Verified Buffer
    if st.session_state.verified_risks_buffer:
        st.markdown("---")
        st.markdown("<div dir='rtl'><h3>🛡️ ریسک‌های تایید شده توسط Checker (آماده انتقال)</h3></div>", unsafe_allow_html=True)
        for idx, item in enumerate(st.session_state.verified_risks_buffer):
            with st.container():
                st.markdown(f"""
                <div class="glass-panel" style="padding: 16px; margin-bottom: 12px; border-right: 4px solid #22c55e;">
                    <strong>عنوان ریسک: {item.get('risk', 'بدون نام')}</strong><br/>
                    <span style="color: #64748b; font-size: 0.95em;">توضیحات پیشنهادی: {item.get('description', '')}</span><br/>
                    <div style="margin-top: 8px;">
                        <span style="background: #e0f2fe; color: #0284c7; padding: 4px 8px; border-radius: 4px; font-weight: bold; font-size: 0.85em;">شدت (S): {item.get('S')}</span>
                        <span style="background: #e0f2fe; color: #0284c7; padding: 4px 8px; border-radius: 4px; font-weight: bold; font-size: 0.85em;">وقوع (O): {item.get('O')}</span>
                        <span style="background: #e0f2fe; color: #0284c7; padding: 4px 8px; border-radius: 4px; font-weight: bold; font-size: 0.85em;">عدم کشف (D): {item.get('D')}</span>
                    </div>
                </div>
                """, unsafe_allow_html=True)
            
            if st.button(f"➕ تایید و انتقال این ریسک به ماتریس", key=f"commit_{idx}"):
                item['RPN'] = int(item['S']) * int(item['O']) * int(item['D'])
                st.session_state.final_assessments.append(item)
                st.session_state.verified_risks_buffer.pop(idx)
                st.rerun()

        if st.button("🧹 پاکسازی بافر ریسک‌های پیشنهادی", type="secondary"):
            st.session_state.verified_risks_buffer = []
            st.rerun()

# --- TAB 2: FMEA Matrix ---
with tab_evaluate:
    st.markdown("<div dir='rtl'><h3>📋 ماتریس FMEA (داده‌های قطعی)</h3></div>", unsafe_allow_html=True)
    if st.session_state.final_assessments:
        df = pd.DataFrame(st.session_state.final_assessments)
        df_sorted = df.sort_values(by="RPN", ascending=False).reset_index(drop=True)
        st.dataframe(df_sorted, use_container_width=True)
        
        if st.button("پاکسازی کامل ماتریس", type="secondary"):
            st.session_state.final_assessments = []
            st.rerun()
    else:
        st.info("هیچ ریسکی تا کنون تایید و به ماتریس منتقل نشده است. ابتدا حلقه ایجنتیک را در تب ۱ اجرا کنید.")

# --- TAB 3: Dashboard ---
with tab_dashboard:
    st.markdown("<div dir='rtl'><h3>📊 مصورسازی متریک‌ها</h3></div>", unsafe_allow_html=True)
    if st.session_state.final_assessments:
        df_dash = pd.DataFrame(st.session_state.final_assessments)
        col1, col2 = st.columns(2)
        with col1:
            fig_bar = px.bar(df_dash, x="risk", y="RPN", color="RPN", color_continuous_scale="Reds", title="اولویت ریسک (RPN)")
            fig_bar.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)')
            st.plotly_chart(fig_bar, use_container_width=True)
        with col2:
            fig_scatter = px.scatter(df_dash, x="O", y="S", size="RPN", color="RPN", title="پراکندگی شدت و وقوع", hover_name="risk")
            fig_scatter.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)')
            st.plotly_chart(fig_scatter, use_container_width=True)
    else:
        st.info("داده‌ای جهت نمایش نمودار وجود ندارد.")

# --- TAB 4: AI & Engine Config ---
with tab_settings:
    st.markdown("<div dir='rtl'><h3>⚙️ پیکربندی مدل و حافظه CoALA</h3></div>", unsafe_allow_html=True)
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("**تنظیمات اتصال Ollama**")
        st.session_state.ollama_url = st.text_input("آدرس API:", value=st.session_state.ollama_url)
        st.session_state.ollama_model = st.text_input("نام مدل لود شده:", value=st.session_state.ollama_model)
    with col2:
        st.markdown("**قوانین حافظه معنایی (Semantic Rules)**")
        rules_text = load_semantic_memory()
        updated_rules = st.text_area("ویرایش فایل `system_rules.md`:", value=rules_text, height=180)
        if st.button("ذخیره قوانین"):
            # ذخیره قوانین در پوشه اختصاصی Ethic_SKILL
            script_dir = os.path.dirname(os.path.abspath(__file__))
            skill_dir = os.path.join(script_dir, "Ethic_SKILL")
            os.makedirs(skill_dir, exist_ok=True)
            file_path = os.path.join(skill_dir, "system_rules.md")
            
            with open(file_path, "w", encoding="utf-8") as f:
                f.write(updated_rules)
            st.success("قوانین با موفقیت در پوشه Ethic_SKILL ذخیره شدند.")